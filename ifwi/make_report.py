# Build a Word report (IFWI_Report.docx) with all figures + narrative.
# Run on ModelWhale AFTER running SIREN Cells 1-3 and U-Net UCELL 1 in the same
# kernel (needs vp_tensor, vi_tensor, wavelet, t, forward_rnn, v_lin, nz, nx, dz,
# extent, ns, NORM_MEAN, NORM_STD). Reads saved predictions from Data/.

import os, glob, subprocess, sys
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", "python-docx"])
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT = "/home/mw/project/IFWI_Project"
CKPT = os.path.join(PROJECT, "checkpoints")
FIGS = os.path.join(PROJECT, "Data", "report_figs")
os.makedirs(FIGS, exist_ok=True)

true_km = vp_tensor.squeeze().cpu().numpy() / 1000
init_km = vi_tensor.squeeze().cpu().numpy() / 1000
lin_km = v_lin / 1000
siren_km = np.load(os.path.join(PROJECT, "Data/siren_vpred_half_best.npy"))
unet_km = np.load(os.path.join(PROJECT, "Data/unet_vpred_half_best.npy"))

def save(fig, name):
    p = os.path.join(FIGS, name)
    fig.savefig(p, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return p

# Fig 1 — true + initial models with shots
fig, ax = plt.subplots(1, 2, figsize=(13, 4))
im = ax[0].imshow(true_km, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
ax[0].plot(np.arange(10, nx - 5, 10) * dz / 1000, np.full(ns, 0.02), "kv", ms=7)
ax[0].set(title="True model (+ 13 shots)", xlabel="Distance (km)", ylabel="Depth (km)")
plt.colorbar(im, ax=ax[0], label="km/s")
im = ax[1].imshow(init_km, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
ax[1].set(title="SIREN initial model (smooth)", xlabel="Distance (km)")
plt.colorbar(im, ax=ax[1], label="km/s")
fig.tight_layout(); f_models = save(fig, "fig1_models.png")

# Fig 2 — linear initial model (U-Net input)
fig, ax = plt.subplots(figsize=(6.5, 4))
im = ax.imshow(lin_km, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
ax.set(title="U-Net initial model (linear)", xlabel="Distance (km)", ylabel="Depth (km)")
plt.colorbar(im, ax=ax, label="km/s")
fig.tight_layout(); f_lin = save(fig, "fig2_linear.png")

# Fig 3 — wavelet + spectrum
w = wavelet.cpu().numpy(); nt = len(w); dt = float(t[1] - t[0])
W = np.abs(np.fft.rfft(w)); fr = np.fft.rfftfreq(nt, dt)
fig, ax = plt.subplots(1, 2, figsize=(13, 4))
ax[0].plot(t.numpy(), w); ax[0].set(title="8 Hz Ricker wavelet", xlabel="Time (s)", ylabel="Amplitude"); ax[0].grid(alpha=0.3)
ax[1].plot(fr, W); ax[1].axvline(8, color="r", ls="--"); ax[1].set(title="Amplitude spectrum", xlabel="Frequency (Hz)", xlim=(0, 40)); ax[1].grid(alpha=0.3)
fig.tight_layout(); f_wav = save(fig, "fig3_wavelet.png")

# Fig 4 — example shot gathers
from generator import wGenerator
wav8 = wGenerator(t, 8.0).ricker().to(device)
import torch
with torch.no_grad():
    _, _, sh, _ = forward_rnn(vmodel=vp_tensor, segment_wavelet=wav8)
sg = sh.squeeze(0).cpu().numpy()
fig, ax = plt.subplots(1, 3, figsize=(15, 4))
for k, s in enumerate([0, ns // 2, ns - 1]):
    clip = np.percentile(np.abs(sg[s]), 99)
    ax[k].imshow(sg[s], aspect="auto", cmap="gray", vmin=-clip, vmax=clip, extent=[0, nx * dz / 1000, nt * dt, 0])
    ax[k].set(title="Shot {}".format(s), xlabel="Receiver x (km)", ylabel="Time (s)" if k == 0 else "")
fig.tight_layout(); f_shots = save(fig, "fig4_shots.png")

# Fig 5/6 — 3-panel results
def panel(pred, name, fname):
    err = np.abs(true_km - pred)
    fig, ax = plt.subplots(1, 3, figsize=(15, 4))
    im = ax[0].imshow(true_km, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
    ax[0].set(title="(a) True model", xlabel="Distance (km)", ylabel="Depth (km)"); plt.colorbar(im, ax=ax[0], label="km/s")
    im = ax[1].imshow(pred, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
    ax[1].set(title="(b) {} prediction".format(name), xlabel="Distance (km)"); plt.colorbar(im, ax=ax[1], label="km/s")
    im = ax[2].imshow(err, vmin=0, vmax=err.max(), extent=extent, aspect=1, cmap="hot_r")
    ax[2].set(title="(c) Pointwise |error|", xlabel="Distance (km)"); plt.colorbar(im, ax=ax[2], label="km/s")
    fig.tight_layout(); return save(fig, fname)
f_siren = panel(siren_km, "SIREN-IFWI", "fig5_siren.png")
f_unet = panel(unet_km, "U-Net-IFWI", "fig6_unet.png")

# Fig 7 — comparison 2x3
fig, ax = plt.subplots(2, 3, figsize=(15, 8))
for row, (name, pred) in enumerate([("SIREN", siren_km), ("U-Net", unet_km)]):
    err = np.abs(true_km - pred)
    im = ax[row, 0].imshow(true_km, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
    ax[row, 0].set(title="(a) True", ylabel="{}\nDepth (km)".format(name)); plt.colorbar(im, ax=ax[row, 0], label="km/s")
    im = ax[row, 1].imshow(pred, vmin=1.5, vmax=5.5, extent=extent, aspect=1, cmap="RdBu_r")
    ax[row, 1].set(title="(b) {} prediction".format(name)); plt.colorbar(im, ax=ax[row, 1], label="km/s")
    im = ax[row, 2].imshow(err, vmin=0, vmax=2.5, extent=extent, aspect=1, cmap="hot_r")
    ax[row, 2].set(title="(c) |error|"); plt.colorbar(im, ax=ax[row, 2], label="km/s")
for a in ax[1, :]:
    a.set_xlabel("Distance (km)")
fig.tight_layout(); f_cmp = save(fig, "fig7_comparison.png")

# Fig 8 — loss curves (use saved PNG if present)
f_loss = os.path.join(CKPT, "SIREN_vs_UNET_loss.png")
f_loss = f_loss if os.path.exists(f_loss) else None

# metrics
def met(p):
    e = np.abs(true_km - p)
    return e.mean(), np.sqrt(((true_km - p) ** 2).mean()), p.min(), p.max()
ms, mu = met(siren_km), met(unet_km)

# ---- build the document ----
doc = Document()
def H(t, lvl=1): doc.add_heading(t, level=lvl)
def P(t):
    p = doc.add_paragraph(t); return p
def IMG(path, w=6.2, cap=None):
    if path and os.path.exists(path):
        doc.add_picture(path, width=Inches(w))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cap:
            c = doc.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.runs[0].italic = True; c.runs[0].font.size = Pt(9)

title = doc.add_heading("Implicit Full Waveform Inversion on Half Marmousi:\nSIREN vs U-Net", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Comparison of two velocity parameterizations for IFWI")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

H("1. Objective")
P("Reproduce implicit full waveform inversion (IFWI) on the left half of the "
  "Marmousi model and compare two velocity parameterizations: a SIREN "
  "(coordinate-based sinusoidal network) and a U-Net (image-based convolutional "
  "network). Both networks are plugged into the same finite-difference wave "
  "simulator and trained by backpropagating the seismic data misfit.")

H("2. Model and acquisition")
P("Half Marmousi, 94 x 144 grid, dz = 15 m. Thirteen surface shots at 150 m "
  "spacing, 144 surface receivers. The true model contains a deep high-velocity "
  "body (~5.5 km/s) at ~1.1 km depth. SIREN starts from a smooth Marmousi "
  "initial model; U-Net takes a linear (depth-increasing) initial model as its "
  "input image, per the supervisor's specification.")
IMG(f_models, cap="Figure 1. True model with shot positions (left) and the smooth SIREN initial model (right).")
IMG(f_lin, w=4.5, cap="Figure 2. Linear initial model used as the U-Net input.")

H("3. Source wavelet and numerical stability")
P("An 8 Hz Ricker wavelet is used (dt = 0.0019 s, 1000 time steps). Numerical "
  "checks pass: stability vmax*dt/dz = 0.70 (< 0.71) and dispersion "
  "vmin/(f*dz) = 12.5 (> 5).")
IMG(f_wav, cap="Figure 3. Ricker wavelet (left) and its amplitude spectrum, peaking at 8 Hz (right).")

H("4. Observed data")
P("Forward modeling on the true model produces the observed shot gathers that "
  "the inversion matches. The direct wave forms a V centered under each shot; "
  "curved reflections below it encode the subsurface layering.")
IMG(f_shots, cap="Figure 4. Example observed shot gathers (left edge, centre, right edge).")

H("5. Method")
P("IFWI parameterizes the velocity field with a neural network. Each epoch the "
  "network outputs a velocity model, the wave simulator generates synthetic "
  "shots, and the mean-squared misfit to the observed shots is backpropagated "
  "into the network weights (Adam optimizer). Both networks are first pretrained "
  "to reproduce their initial model, then inverted. 5000 epochs per method.")

H("6. Cycle-skipping and frequency continuation")
P("Single-frequency 8 Hz inversion cycle-skipped: the data misfit dropped ~99% "
  "but the model error increased (-14.8%), with the deep body capped near "
  "4.0 km/s. The fix is frequency continuation - invert from low to high "
  "frequency (3 -> 5 -> 8 Hz, 1500 + 1500 + 2000 epochs). The long wavelength "
  "at 3 Hz recovers the large-scale structure without cycle-skipping; higher "
  "frequencies then add detail.")

H("7. SIREN result")
IMG(f_siren, cap="Figure 5. SIREN-IFWI: true model, prediction, and pointwise absolute error.")
P("Final MAE = {:.4f} km/s, RMS = {:.4f} km/s. Predicted range {:.3f}-{:.3f} km/s. "
  "SIREN recovers the shallow and mid-section structure but leaves the deep "
  "high-velocity body capped near 4.0 km/s.".format(ms[0], ms[1], ms[2], ms[3]))

H("8. U-Net result")
IMG(f_unet, cap="Figure 6. U-Net-IFWI: true model, prediction, and pointwise absolute error.")
P("Final MAE = {:.4f} km/s, RMS = {:.4f} km/s. Predicted range {:.3f}-{:.3f} km/s. "
  "U-Net recovers the deep high-velocity body that SIREN missed (reaching "
  "{:.2f} km/s, a slight overshoot of the 5.5 km/s truth).".format(mu[0], mu[1], mu[2], mu[3], mu[3]))

H("9. Comparison")
IMG(f_cmp, cap="Figure 7. Side-by-side comparison: SIREN (top) and U-Net (bottom).")
t = doc.add_table(rows=5, cols=3); t.style = "Light Grid Accent 1"
rows = [["Metric", "SIREN", "U-Net"],
        ["Final MAE (km/s)", "{:.4f}".format(ms[0]), "{:.4f}".format(mu[0])],
        ["Final RMS (km/s)", "{:.4f}".format(ms[1]), "{:.4f}".format(mu[1])],
        ["Predicted vmax (km/s)", "{:.3f}".format(ms[3]), "{:.3f}".format(mu[3])],
        ["Parameters", "~0.05 M", "~7.76 M"]]
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        t.cell(i, j).text = v
if f_loss:
    IMG(f_loss, cap="Figure 8. Data-misfit convergence for both methods (log scale). Jumps mark the 3->5->8 Hz stage transitions.")
P("U-Net outperforms SIREN on this half-aperture problem in every model-domain "
  "metric, and crucially recovers the deep body. Note that SIREN reaches a lower "
  "data misfit at 3 Hz yet yields a worse model - low data misfit does not "
  "guarantee a correct model. U-Net's convolutional structure acts as a stronger "
  "model-space prior, at ~150x more parameters than SIREN.")

H("10. Conclusions")
P("Frequency continuation is essential for both methods on the half Marmousi "
  "model. With it, U-Net (MAE 0.20, RMS 0.35 km/s) clearly outperforms SIREN "
  "(MAE 0.31, RMS 0.61 km/s) and recovers the deep high-velocity structure. The "
  "remaining error concentrates in the deep section, limited by the surface-only, "
  "half-aperture acquisition geometry shared by both methods.")

out = os.path.join(PROJECT, "IFWI_Report.docx")
doc.save(out)
print("Report written to:", out)
print("Figures in:", FIGS)
